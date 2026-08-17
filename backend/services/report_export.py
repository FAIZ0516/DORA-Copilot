"""Render a report to PDF, DOCX or CSV.

PDF is produced server-side with ReportLab rather than browser print, so the
output is identical regardless of who exports it and page furniture (numbers,
version, repeated table headers) is under our control. ReportLab is a pure
Python wheel -- WeasyPrint would drag GTK/Cairo system libraries in, which is a
poor trade on the Windows dev environment this project targets.

Layout is deterministic here. Nothing in this module calls a model.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from typing import Any

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.graphics.shapes import Circle, Drawing, String
from .report_branding import (
    FOOTER_COVER_BOTTOM,
    FOOTER_COVER_HEIGHT,
    FOOTER_COVER_WIDTH,
    FOOTER_NUMBER_BASELINE,
    TEMPLATE_MARGIN_BOTTOM,
    TEMPLATE_MARGIN_LEFT,
    TEMPLATE_MARGIN_RIGHT,
    TEMPLATE_MARGIN_TOP,
    apply_template,
    template_available,
)
from .report_charts import describe, draw_chart, truncation_note
from reportlab.platypus import (
    BaseDocTemplate,
    Flowable,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

# Brand palette, aligned with the app's chart tokens so an exported report and
# the on-screen dashboard read as one product.
INK = colors.HexColor("#1f2d3d")
MUTED = colors.HexColor("#5b6b80")
ACCENT = colors.HexColor("#0b68b5")
RULE = colors.HexColor("#dbe3ee")
SURFACE = colors.HexColor("#f6f8fc")
RISK = colors.HexColor("#b4532a")

PAGE_WIDTH, PAGE_HEIGHT = A4
MARGIN = 18 * mm

def safe_filename(title: str, extension: str) -> str:
    """A predictable, filesystem-safe download name."""

    cleaned = re.sub(r"[^\w\s-]", "", title or "report").strip()
    cleaned = re.sub(r"[\s_]+", "-", cleaned).strip("-").lower() or "report"
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    return f"{cleaned[:60]}-{stamp}.{extension}"


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "cover_title": ParagraphStyle(
            "cover_title", parent=base["Title"], fontName="Helvetica-Bold",
            fontSize=27, leading=32, textColor=INK, spaceAfter=6,
        ),
        "cover_meta": ParagraphStyle(
            "cover_meta", parent=base["Normal"], fontName="Helvetica",
            fontSize=10, leading=16, textColor=MUTED,
        ),
        "heading": ParagraphStyle(
            "heading", parent=base["Heading2"], fontName="Helvetica-Bold",
            fontSize=13, leading=17, textColor=INK, spaceBefore=14, spaceAfter=5,
        ),
        "body": ParagraphStyle(
            "body", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5,
            leading=14.5, textColor=INK, alignment=TA_LEFT, spaceAfter=6,
        ),
        "label": ParagraphStyle(
            "label", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=6.5,
            leading=9, textColor=MUTED, spaceAfter=3,
        ),
        "note": ParagraphStyle(
            "note", parent=base["Normal"], fontName="Helvetica-Oblique", fontSize=8,
            leading=12, textColor=MUTED, spaceAfter=4,
        ),
        "cell": ParagraphStyle(
            "cell", parent=base["Normal"], fontName="Helvetica", fontSize=7.6, leading=10,
            textColor=INK,
        ),
        "chart_caption": ParagraphStyle(
            "chart_caption", parent=base["Normal"], fontName="Helvetica", fontSize=8.2,
            leading=12.5, textColor=MUTED, spaceBefore=4, spaceAfter=6,
        ),
        "cell_head": ParagraphStyle(
            "cell_head", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=7.6,
            leading=10, textColor=colors.white,
        ),
        "weekly_eyebrow": ParagraphStyle(
            "weekly_eyebrow", parent=base["Normal"], fontName="Helvetica-Bold",
            fontSize=9, leading=11, textColor=ACCENT, spaceAfter=5,
        ),
        "weekly_title": ParagraphStyle(
            "weekly_title", parent=base["Title"], fontName="Helvetica-Bold",
            fontSize=21, leading=24, textColor=INK, alignment=TA_LEFT, spaceAfter=5,
        ),
        "weekly_section": ParagraphStyle(
            "weekly_section", parent=base["Heading2"], fontName="Helvetica-Bold",
            fontSize=9.2, leading=12, textColor=ACCENT, spaceBefore=8, spaceAfter=6,
        ),
        "weekly_body": ParagraphStyle(
            "weekly_body", parent=base["BodyText"], fontName="Helvetica",
            fontSize=8.4, leading=11.8, textColor=INK, spaceAfter=4,
        ),
        "weekly_meta": ParagraphStyle(
            "weekly_meta", parent=base["Normal"], fontName="Helvetica",
            fontSize=8, leading=11, textColor=MUTED, spaceAfter=5,
        ),
        "weekly_kpi_label": ParagraphStyle(
            "weekly_kpi_label", parent=base["Normal"], fontName="Helvetica-Bold",
            fontSize=7, leading=9, textColor=MUTED,
        ),
        "weekly_kpi_value": ParagraphStyle(
            "weekly_kpi_value", parent=base["Normal"], fontName="Helvetica-Bold",
            fontSize=17, leading=20, textColor=INK,
        ),
        "weekly_kpi_note": ParagraphStyle(
            "weekly_kpi_note", parent=base["Normal"], fontName="Helvetica",
            fontSize=6.6, leading=8, textColor=MUTED,
        ),
        "weekly_number": ParagraphStyle(
            "weekly_number", parent=base["Normal"], fontName="Helvetica-Bold",
            fontSize=8, leading=10, alignment=TA_CENTER, textColor=colors.white,
        ),
        "weekly_risk_number": ParagraphStyle(
            "weekly_risk_number", parent=base["Normal"], fontName="Helvetica-Bold",
            fontSize=7.5, leading=10, textColor=RISK,
        ),
        "weekly_item_title": ParagraphStyle(
            "weekly_item_title", parent=base["Normal"], fontName="Helvetica-Bold",
            fontSize=9, leading=11, textColor=INK, spaceAfter=2,
        ),
        "weekly_table_head": ParagraphStyle(
            "weekly_table_head", parent=base["Normal"], fontName="Helvetica-Bold",
            fontSize=7.2, leading=9, textColor=ACCENT,
        ),
    }


def _escape(text: Any) -> str:
    """Escape for ReportLab's mini-HTML so user text cannot break layout."""

    return (
        str(text if text is not None else "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _scope_line(scope: dict[str, Any]) -> str:
    parts = [
        f"Project {scope.get('project')}" if scope.get("project") else None,
        f"Squad {scope.get('squad')}" if scope.get("squad") else "All squads",
        f"Sprint {scope.get('sprint')}" if scope.get("sprint") else None,
        f"Release {scope.get('release')}" if scope.get("release") else None,
    ]
    period = None
    if scope.get("date_from") or scope.get("date_to"):
        period = f"{scope.get('date_from') or 'start'} to {scope.get('date_to') or 'today'}"
    parts.append(f"Period {period}" if period else "All available dates")
    return " · ".join(part for part in parts if part)


def _uses_delivery_layout(report: dict[str, Any]) -> bool:
    return report.get("template") in {"weekly_scrum", "executive_summary"}


class _ReportDoc(BaseDocTemplate):
    """Adds the corporate furniture and running page number."""

    def __init__(self, buffer: io.BytesIO, *, report: dict[str, Any]) -> None:
        branded = template_available()
        # Inside the template's frame when branded, so body text never runs
        # under the masthead or across the footer rule.
        left = TEMPLATE_MARGIN_LEFT if branded else MARGIN
        right = TEMPLATE_MARGIN_RIGHT if branded else MARGIN
        top = TEMPLATE_MARGIN_TOP if branded else MARGIN
        bottom = TEMPLATE_MARGIN_BOTTOM if branded else MARGIN + 8 * mm
        super().__init__(
            buffer, pagesize=A4,
            leftMargin=left, rightMargin=right, topMargin=top, bottomMargin=bottom,
            title=report.get("title") or "Report", author="ZARA Report Studio",
        )
        self._report = report
        self._branded = branded
        self._left = left
        self._right = right
        frame = Frame(left, bottom, PAGE_WIDTH - left - right,
                      PAGE_HEIGHT - top - bottom, id="body")
        self.addPageTemplates([
            PageTemplate(id="cover", frames=[frame], onPage=self._footer),
            PageTemplate(id="body", frames=[frame], onPage=self._footer),
        ])

    def _footer(self, canvas, doc) -> None:  # noqa: ANN001 - ReportLab signature
        canvas.saveState()
        if self._branded:
            # Keep the supplied corporate masthead and footer, but mask its
            # decorative body watermark so all exported reports have the clean
            # white management-report surface used by the weekly reference.
            mask_left = TEMPLATE_MARGIN_LEFT - 5 * mm
            mask_bottom = TEMPLATE_MARGIN_BOTTOM - 3 * mm
            canvas.setFillColor(colors.white)
            canvas.rect(
                mask_left,
                mask_bottom,
                PAGE_WIDTH - (2 * mask_left),
                PAGE_HEIGHT - TEMPLATE_MARGIN_TOP - mask_bottom,
                stroke=0,
                fill=1,
            )
            # The template carries its own rule, and a static "1" that would
            # otherwise repeat on every page. Cover it, then write the real
            # number. Content is merged above the template, so this hides it.
            canvas.setFillColor(colors.white)
            canvas.rect(
                PAGE_WIDTH - self._right - FOOTER_COVER_WIDTH, FOOTER_COVER_BOTTOM,
                FOOTER_COVER_WIDTH, FOOTER_COVER_HEIGHT, stroke=0, fill=1,
            )
            baseline = FOOTER_NUMBER_BASELINE
        else:
            canvas.setStrokeColor(RULE)
            canvas.setLineWidth(0.5)
            canvas.line(self._left, MARGIN + 6 * mm, PAGE_WIDTH - self._right, MARGIN + 6 * mm)
            baseline = MARGIN + 2 * mm
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawRightString(
            PAGE_WIDTH - self._right, baseline, f"Page {canvas.getPageNumber()}"
        )
        canvas.restoreState()


def _kpi_table(payload: dict[str, Any], styles: dict[str, ParagraphStyle]) -> Table | None:
    items = [item for item in (payload.get("items") or []) if isinstance(item, dict)]
    if not items:
        return None
    cells = [
        [
            Paragraph(_escape(item.get("label", "")), styles["label"]),
            Paragraph(f"<b>{_escape(item.get('value', ''))}</b>", styles["body"]),
        ]
        for item in items[:8]
    ]
    # Two KPIs per row keeps values large enough to scan.
    rows = [sum(cells[i : i + 2], []) for i in range(0, len(cells), 2)]
    width = (PAGE_WIDTH - 2 * MARGIN) / 4
    table = Table(rows, colWidths=[width] * 4, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), SURFACE),
        ("BOX", (0, 0), (-1, -1), 0.5, RULE),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, RULE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    return table


def _data_table(
    payload: dict[str, Any], styles: dict[str, ParagraphStyle], *, content_width: float | None = None
) -> list[Any]:
    columns = payload.get("columns") or []
    rows = payload.get("rows") or []
    if not columns or not rows:
        return [Paragraph("No rows were available for this table.", styles["note"])]
    keys = [column.get("key") for column in columns if isinstance(column, dict)]
    headers = [
        Paragraph(_escape(column.get("label") or column.get("key")), styles["cell_head"])
        for column in columns
        if isinstance(column, dict)
    ]
    # Cap rows so one table cannot produce a hundred-page appendix.
    limit = 60
    body = [
        [Paragraph(_escape(row.get(key)), styles["cell"]) for key in keys]
        for row in rows[:limit]
    ]
    width = (content_width or (PAGE_WIDTH - 2 * MARGIN)) / max(len(keys), 1)
    # repeatRows=1 reprints the header on every page the table spills onto.
    table = Table([headers, *body], colWidths=[width] * len(keys), repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, SURFACE]),
        ("GRID", (0, 0), (-1, -1), 0.4, RULE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    flow: list[Any] = [table]
    if len(rows) > limit:
        flow.append(Paragraph(
            f"Showing the first {limit} of {len(rows)} rows. Export the table as CSV "
            f"for the complete set.", styles["note"],
        ))
    return flow


def _chart_block(payload: dict[str, Any], styles: dict[str, ParagraphStyle]) -> list[Any]:
    """A real vector chart, with a factual reading of it underneath.

    Charts used to export as a table of values. Accurate, but a management
    report is read at a glance, so the numbers are now drawn. Vector output
    stays sharp in print and needs no headless browser on the server. The
    caption states only what the data already says, so it cannot drift from
    the plot beside it.
    """

    flow: list[Any] = []
    drawing = draw_chart(payload, PAGE_WIDTH - 2 * MARGIN)
    if drawing is None:
        return [Paragraph("No chart data was available for this section.", styles["note"])]

    flow.append(drawing)
    note = truncation_note(payload)
    if note:
        flow.append(Paragraph(_escape(note), styles["note"]))
    reading = describe(payload)
    if reading:
        flow.append(Paragraph(_escape(reading), styles["chart_caption"]))
    return flow


def _report_intro(report: dict[str, Any], styles: dict[str, ParagraphStyle]) -> list[Any]:
    """Compact first-page context for non-delivery templates; never a cover."""

    scope = report.get("scope") or {}
    return [
        Paragraph("ZARA DELIVERY REPORT", styles["label"]),
        Paragraph(_escape(report.get("title") or "Report"), styles["cover_title"]),
        Paragraph(_escape(_scope_line(scope)), styles["cover_meta"]),
        Paragraph(
            _escape(f"Data as of {_display_day(report.get('data_as_of'))}"),
            styles["cover_meta"],
        ),
        Spacer(1, 5 * mm),
    ]


def _display_day(value: Any) -> str:
    if not value:
        return "Not recorded"
    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except ValueError:
            return str(value)
    return parsed.strftime("%d %b %Y")


def _content_lines(section: dict[str, Any]) -> list[str]:
    return [
        line.strip().lstrip("-• ").strip()
        for line in str(section.get("content") or "").splitlines()
        if line.strip()
    ]


class _WeeklyKpiGrid(Flowable):
    """Two-by-two rounded KPI cards populated from verified payload values."""

    def __init__(self, cards: list[dict[str, Any]], width: float) -> None:
        super().__init__()
        self.cards = cards
        self.width = width
        self.height = 126

    def wrap(self, _available_width: float, _available_height: float) -> tuple[float, float]:
        return self.width, self.height

    def draw(self) -> None:
        canvas = self.canv
        gap = 9
        card_width = (self.width - gap) / 2
        card_height = (self.height - gap) / 2
        for index, card in enumerate(self.cards):
            column = index % 2
            row = index // 2
            x = column * (card_width + gap)
            y = self.height - (row + 1) * card_height - row * gap
            canvas.setFillColor(colors.white)
            canvas.setStrokeColor(RULE)
            canvas.setLineWidth(0.7)
            canvas.roundRect(x, y, card_width, card_height, 6, stroke=1, fill=1)
            canvas.setStrokeColor(card["accent"])
            canvas.setLineWidth(3)
            canvas.line(x + 7, y + card_height - 2, x + card_width - 7, y + card_height - 2)
            canvas.setFillColor(MUTED)
            canvas.setFont("Helvetica-Bold", 7)
            canvas.drawString(x + 10, y + card_height - 17, card["label"])
            canvas.setFillColor(INK)
            canvas.setFont("Helvetica-Bold", 17)
            canvas.drawString(x + 10, y + card_height - 38, str(card["value"]))
            canvas.setFillColor(MUTED)
            canvas.setFont("Helvetica", 6.7)
            canvas.drawString(x + 10, y + 8, card["note"])


def _weekly_kpi_grid(
    payload: dict[str, Any], _styles: dict[str, ParagraphStyle]
) -> Flowable | None:
    items = {
        str(item.get("key")): item
        for item in (payload.get("items") or [])
        if isinstance(item, dict)
    }
    ordered = ["completion_pct", "active_work", "impeded_work", "open_bugs"]
    if not any(key in items for key in ordered):
        return None

    completed = items.get("completed_work", {}).get("value", "Unavailable")
    total = items.get("total_work", {}).get("value", "Unavailable")
    notes = {
        "completion_pct": f"{completed} of {total} scoped tickets",
        "active_work": "Unresolved items",
        "impeded_work": "Items currently impeded",
        "open_bugs": "Unresolved bug tickets",
    }
    accents = {
        "completion_pct": ACCENT,
        "active_work": ACCENT,
        "impeded_work": colors.HexColor("#d17a16"),
        "open_bugs": colors.HexColor("#ef4654"),
    }
    labels = {
        "completion_pct": "SPRINT COMPLETION",
        "active_work": "OPEN WORK",
        "impeded_work": "ACTIVE BLOCKERS",
        "open_bugs": "OPEN BUGS",
    }
    cards: list[dict[str, Any]] = []
    for key in ordered:
        item = items.get(key, {})
        cards.append({
            "label": labels[key],
            "value": item.get("value", "Unavailable"),
            "note": notes[key],
            "accent": accents[key],
        })
    return _WeeklyKpiGrid(
        cards, PAGE_WIDTH - TEMPLATE_MARGIN_LEFT - TEMPLATE_MARGIN_RIGHT
    )


def _weekly_feature_table(
    payload: dict[str, Any], styles: dict[str, ParagraphStyle]
) -> list[Any]:
    rows = payload.get("rows") or []
    if not rows:
        return [Paragraph("No Feature rows were available for this scope.", styles["note"])]

    def status_colour(value: Any) -> str:
        status = str(value or "").casefold()
        if any(token in status for token in ("done", "closed", "released", "complete")):
            return "#16835b"
        if any(token in status for token in ("blocked", "impeded", "cancelled", "rejected")):
            return "#c2414b"
        if any(token in status for token in ("progress", "test", "review", "development")):
            return "#b66a10"
        return "#4c5d73"

    content_width = PAGE_WIDTH - TEMPLATE_MARGIN_LEFT - TEMPLATE_MARGIN_RIGHT
    table_rows: list[list[Any]] = [[
        Paragraph("FEATURE ID", styles["weekly_table_head"]),
        Paragraph("FEATURE NAME", styles["weekly_table_head"]),
        Paragraph("STATUS", styles["weekly_table_head"]),
    ]]
    for row in rows:
        status = row.get("status") or "Needs Input"
        table_rows.append([
            Paragraph(f"<b>{_escape(row.get('feature'))}</b>", styles["cell"]),
            Paragraph(_escape(row.get("feature_name")), styles["cell"]),
            Paragraph(
                f'<font color="{status_colour(status)}"><b>{_escape(status)}</b></font>',
                styles["cell"],
            ),
        ])
    table = Table(
        table_rows,
        colWidths=[content_width * .21, content_width * .54, content_width * .25],
        repeatRows=1,
        hAlign="LEFT",
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef5fb")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fbfcfe")]),
        ("GRID", (0, 0), (-1, -1), 0.4, RULE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return [table]


def _risk_title(text: str) -> str:
    lowered = text.casefold()
    if "block" in lowered or "impeded" in lowered:
        return "Blocked work requires follow-up"
    if "bug" in lowered:
        return "Open bugs remain in the current scope"
    if "oldest" in lowered or "age" in lowered:
        return "Ageing work needs review"
    if "assign" in lowered or "owner" in lowered:
        return "Open work needs ownership"
    if "completion" in lowered:
        return "Completion needs attention"
    if "status category" in lowered:
        return "Status data requires review"
    return "Delivery condition requires attention"


def _action_title(text: str) -> str:
    lowered = text.casefold()
    if "block" in lowered or "impeded" in lowered:
        return "Resolve blocker ownership"
    if "bug" in lowered:
        return "Prioritise unresolved bugs"
    if "age" in lowered or "oldest" in lowered or "long-running" in lowered:
        return "Review ageing work"
    if "owner" in lowered or "unassigned" in lowered:
        return "Confirm work ownership"
    if "remaining" in lowered or "priorit" in lowered:
        return "Prioritise remaining work"
    if "missing" in lowered or "status" in lowered:
        return "Resolve reporting gaps"
    return "Continue delivery monitoring"


def _weekly_risk_blocks(
    section: dict[str, Any], styles: dict[str, ParagraphStyle]
) -> list[Any]:
    lines = _content_lines(section) or ["No verified delivery risk currently requires attention in this scope."]
    if all(line.casefold().startswith("no verified delivery risk") for line in lines):
        return [Paragraph(_escape(lines[0]), styles["weekly_body"])]
    flow: list[Any] = []
    content_width = PAGE_WIDTH - TEMPLATE_MARGIN_LEFT - TEMPLATE_MARGIN_RIGHT
    for index, line in enumerate(lines, start=1):
        card = Table(
            [[Paragraph(f"{index:02d}", styles["weekly_risk_number"]),
              [Paragraph(_escape(_risk_title(line)), styles["weekly_item_title"]),
               Paragraph(_escape(line), styles["weekly_meta"])]]],
            colWidths=[30, content_width - 30], hAlign="LEFT",
        )
        card.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), .7, RULE),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fcfdff")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        flow.extend([KeepTogether([card]), Spacer(1, 5 * mm)])
    return flow


def _weekly_action_blocks(
    section: dict[str, Any], styles: dict[str, ParagraphStyle]
) -> list[Any]:
    lines = _content_lines(section) or ["Continue monitoring the current delivery scope."]
    if all(line.casefold().startswith("no new action is required") for line in lines):
        return [Paragraph(_escape(lines[0]), styles["weekly_body"])]
    flow: list[Any] = []
    content_width = PAGE_WIDTH - TEMPLATE_MARGIN_LEFT - TEMPLATE_MARGIN_RIGHT
    for index, line in enumerate(lines, start=1):
        number = Drawing(24, 24)
        number.add(Circle(12, 12, 12, fillColor=ACCENT, strokeColor=ACCENT))
        number.add(String(
            12, 9, str(index), textAnchor="middle", fillColor=colors.white,
            fontName="Helvetica-Bold", fontSize=8,
        ))
        row = Table(
            [[number, [Paragraph(_escape(_action_title(line)), styles["weekly_item_title"]),
                       Paragraph(_escape(line), styles["weekly_meta"])]]],
            colWidths=[34, content_width - 34], hAlign="LEFT",
        )
        row.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        flow.extend([KeepTogether([row]), Spacer(1, 4 * mm)])
    return flow


def _delivery_context(
    report: dict[str, Any], styles: dict[str, ParagraphStyle], *, insights: bool = False
) -> list[Any]:
    scope = report.get("scope") or {}
    squad = scope.get("squad") or "All DCP Squads"
    sprint = scope.get("sprint") or "All Sprints"
    project = scope.get("project") or "DCPM"
    title = f"{squad} - Delivery Insights" if insights else f"{squad} - {sprint}"
    meta = (
        f"Project {project}  |  Squad {squad}  |  {sprint}"
        if insights
        else f"Project {project}  |  Data as of {_display_day(report.get('data_as_of'))}"
    )
    return [
        Paragraph("ZARA WEEKLY DELIVERY REPORT", styles["weekly_eyebrow"]),
        Paragraph(_escape(title), styles["weekly_title"]),
        Paragraph(_escape(meta), styles["weekly_meta"]),
        Spacer(1, 3 * mm),
    ]


def _delivery_section_block(
    section: dict[str, Any], styles: dict[str, ParagraphStyle]
) -> list[Any]:
    type_ = section.get("type")
    title = section.get("title") or type_.replace("_", " ").title()
    payload = section.get("payload") or {}
    content = str(section.get("content") or "").strip()
    block: list[Any] = [
        Paragraph(_escape(title.upper()), styles["weekly_section"]),
    ]

    if type_ == "kpi_group":
        grid = _weekly_kpi_grid(payload, styles)
        block.append(
            grid
            if grid is not None
            else Paragraph("Verified KPI values are not available for this scope.", styles["note"])
        )
    elif type_ == "feature_status":
        block.extend(_weekly_feature_table(payload, styles))
    elif type_ == "key_finding":
        lines = _content_lines(section)
        block.extend(
            Paragraph(_escape(f"- {line}"), styles["weekly_body"])
            for line in lines
        )
    elif type_ == "risk":
        block.extend(_weekly_risk_blocks(section, styles))
    elif type_ in {"recommendation", "action_list"}:
        block.extend(_weekly_action_blocks(section, styles))
    elif type_ in {"data_table"}:
        block.extend(_data_table(
            payload, styles,
            content_width=PAGE_WIDTH - TEMPLATE_MARGIN_LEFT - TEMPLATE_MARGIN_RIGHT,
        ))
    elif type_ == "chart":
        block.extend(_chart_block(payload, styles))
    elif content:
        block.extend(
            Paragraph(_escape(paragraph), styles["weekly_body"])
            for paragraph in content.split("\n")
            if paragraph.strip()
        )
    else:
        block.append(Paragraph("This section has not been written yet.", styles["note"]))
    block.append(Spacer(1, 3 * mm))
    return block


def _render_delivery_pdf(report: dict[str, Any]) -> bytes:
    """Render the saved delivery document in its current order and visibility."""

    styles = _styles()
    sections = [
        section
        for section in sorted(report.get("sections") or [], key=lambda item: item.get("position", 0))
        if section.get("visible", True) and section.get("type") not in {"cover", "methodology"}
    ]
    flow: list[Any] = _delivery_context(report, styles)
    rendered_count = 0
    insights_started = False
    for section in sections:
        type_ = section.get("type")
        if type_ == "page_break":
            flow.append(PageBreak())
            flow.extend(_delivery_context(report, styles, insights=True))
            insights_started = True
            continue
        if (
            not insights_started
            and rendered_count >= 2
            and type_ in {"risk", "recommendation", "action_list", "data_quality"}
        ):
            flow.append(PageBreak())
            flow.extend(_delivery_context(report, styles, insights=True))
            insights_started = True
        flow.extend(_delivery_section_block(section, styles))
        rendered_count += 1

    buffer = io.BytesIO()
    doc = _ReportDoc(buffer, report=report)
    doc.build(flow)
    return apply_template(buffer.getvalue())


def render_pdf(report: dict[str, Any]) -> bytes:
    """Render the whole report. Returns PDF bytes; never writes to disk."""

    if _uses_delivery_layout(report):
        return _render_delivery_pdf(report)

    styles = _styles()
    buffer = io.BytesIO()
    doc = _ReportDoc(buffer, report=report)
    flow: list[Any] = _report_intro(report, styles)

    for section in sorted(report.get("sections") or [], key=lambda s: s.get("position", 0)):
        if not section.get("visible", True) or section.get("type") == "methodology":
            continue
        type_ = section.get("type")
        if type_ == "cover":
            continue
        if type_ == "page_break":
            flow.append(PageBreak())
            continue

        block: list[Any] = []
        if section.get("title"):
            block.append(Paragraph(_escape(section["title"]), styles["heading"]))

        if section.get("needs_review"):
            block.append(Paragraph(
                "Manually edited after validation — figures in this section have not been "
                "re-verified against the source evidence.", styles["note"],
            ))

        payload = section.get("payload") or {}
        if type_ == "kpi_group":
            table = _kpi_table(payload, styles)
            if table is not None:
                block.append(table)
        elif type_ in {"data_table", "feature_status"}:
            block.extend(_data_table(payload, styles))
        elif type_ == "chart":
            block.extend(_chart_block(payload, styles))
        content = (section.get("content") or "").strip()
        if content:
            for paragraph in [p for p in content.split("\n") if p.strip()]:
                block.append(Paragraph(_escape(paragraph), styles["body"]))
        elif type_ in {"executive_summary", "rich_text", "key_finding"} and not payload:
            block.append(Paragraph("This section has not been written yet.", styles["note"]))

        block.append(Spacer(1, 3 * mm))
        # Keep a heading with at least the start of its content.
        flow.append(KeepTogether(block[:2]) if len(block) > 2 else KeepTogether(block))
        flow.extend(block[2:] if len(block) > 2 else [])

    doc.build(flow)
    return apply_template(buffer.getvalue())


def render_docx(report: dict[str, Any]) -> bytes:
    """Editable Word version, mirroring the PDF's structure."""

    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    document.core_properties.title = report.get("title") or "Report"
    scope = report.get("scope") or {}
    if _uses_delivery_layout(report):
        document.add_heading("ZARA WEEKLY DELIVERY REPORT", level=0)
        document.add_heading(
            f"{scope.get('squad') or 'All DCP Squads'} - {scope.get('sprint') or 'All Sprints'}",
            level=1,
        )
        document.add_paragraph(f"Project {scope.get('project') or 'DCPM'}")
    else:
        document.add_heading(report.get("title") or "Report", level=0)
        document.add_paragraph(_scope_line(scope))
    meta = document.add_paragraph()
    meta.add_run(
        f"Data as of {_display_day(report.get('data_as_of'))}"
    ).font.size = Pt(8)

    for section in sorted(report.get("sections") or [], key=lambda s: s.get("position", 0)):
        if not section.get("visible", True) or section.get("type") in {"cover", "methodology"}:
            continue
        if section.get("type") == "page_break":
            document.add_page_break()
            continue
        if section.get("title"):
            document.add_heading(section["title"], level=1)
        payload = section.get("payload") or {}
        if section.get("type") == "kpi_group":
            items = {
                str(item.get("key")): item
                for item in (payload.get("items") or [])
                if isinstance(item, dict)
            }
            ordered = ["completion_pct", "active_work", "impeded_work", "open_bugs"]
            if any(key in items for key in ordered):
                completed = items.get("completed_work", {}).get("value", "Unavailable")
                total = items.get("total_work", {}).get("value", "Unavailable")
                notes = {
                    "completion_pct": f"{completed} of {total} scoped tickets",
                    "active_work": "Unresolved items",
                    "impeded_work": "Items currently impeded",
                    "open_bugs": "Unresolved bug tickets",
                }
                labels = {
                    "completion_pct": "SPRINT COMPLETION",
                    "active_work": "OPEN WORK",
                    "impeded_work": "ACTIVE BLOCKERS",
                    "open_bugs": "OPEN BUGS",
                }
                table = document.add_table(rows=2, cols=2)
                table.style = "Light Grid Accent 1"
                for index, key in enumerate(ordered):
                    cell = table.cell(index // 2, index % 2)
                    cell.text = ""
                    label_run = cell.paragraphs[0].add_run(labels[key])
                    label_run.bold = True
                    label_run.font.size = Pt(8)
                    label_run.font.color.rgb = RGBColor(91, 107, 128)
                    value_run = cell.add_paragraph().add_run(
                        str(items.get(key, {}).get("value", "Unavailable"))
                    )
                    value_run.bold = True
                    value_run.font.size = Pt(18)
                    note_run = cell.add_paragraph().add_run(notes[key])
                    note_run.font.size = Pt(8)
                    note_run.font.color.rgb = RGBColor(91, 107, 128)
        elif section.get("type") in {"data_table", "feature_status", "chart"}:
            columns, rows = _tabular(payload)
            if columns:
                table = document.add_table(rows=1, cols=len(columns))
                table.style = "Light Grid Accent 1"
                for index, name in enumerate(columns):
                    table.rows[0].cells[index].text = str(name)
                for row in rows[:60]:
                    cells = table.add_row().cells
                    for index, value in enumerate(row):
                        cells[index].text = "" if value is None else str(value)

        for paragraph in [p for p in (section.get("content") or "").split("\n") if p.strip()]:
            document.add_paragraph(paragraph)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _tabular(payload: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    """Flatten a table or chart payload into header + rows for CSV/DOCX."""

    if payload.get("columns"):
        columns = [c for c in payload["columns"] if isinstance(c, dict)]
        keys = [c.get("key") for c in columns]
        headers = [c.get("label") or c.get("key") for c in columns]
        return headers, [[row.get(key) for key in keys] for row in payload.get("rows") or []]
    series = [s for s in (payload.get("series") or []) if isinstance(s, dict)]
    if series:
        x_key = payload.get("x_key") or "label"
        headers = [payload.get("x_label") or x_key] + [s.get("label") for s in series]
        rows = [
            [row.get(x_key)] + [row.get(s.get("key")) for s in series]
            for row in payload.get("data") or []
        ]
        return headers, rows
    return [], []


def render_csv(payload: dict[str, Any]) -> str:
    """One table block as CSV. Full rows, no display cap."""

    headers, rows = _tabular(payload or {})
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    if headers:
        writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if value is None else value for value in row])
    return buffer.getvalue()


__all__ = ["render_csv", "render_docx", "render_pdf", "safe_filename"]
